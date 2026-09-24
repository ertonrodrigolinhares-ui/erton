"""Ferramentas que a IA pode usar sozinha: planilhas, arquivos, código, mensagens, lembretes...

Cada ferramenta é uma função Python comum, com docstring explicando para a IA quando usá-la.
Ações que alteram arquivos existentes ou executam código pedem confirmação do usuário.
Mensagens e e-mails nunca são enviados sozinhos: o Jarvis abre tudo pronto e você aperta enviar.
"""

import functools
import io
import json
import os
import re
import subprocess
import sys
import threading
import urllib.parse
import webbrowser
from pathlib import Path
from typing import Callable

# Pasta onde o Jarvis cria e procura arquivos quando você não diz o caminho completo.
from . import memoria
from .config import PASTA_TRABALHO

PROGRAMAS = {
    "excel": "excel.exe", "word": "winword.exe", "powerpoint": "powerpnt.exe",
    "outlook": "outlook.exe", "calculadora": "calc.exe", "bloco de notas": "notepad.exe",
    "paint": "mspaint.exe", "explorador de arquivos": "explorer.exe", "prompt de comando": "cmd.exe",
}

LIMITE_SAIDA = 4000


def _caminho(arquivo: str, extensao: str = "") -> Path:
    caminho = Path(os.path.expandvars(os.path.expanduser(arquivo.strip().strip('"'))))
    if not caminho.is_absolute():
        caminho = PASTA_TRABALHO / caminho
    if extensao and caminho.suffix.lower() != extensao:
        caminho = caminho.with_suffix(extensao)
    return caminho


def _numero_ou_texto(valor: str):
    """Converte "1.234,56" ou "1234.56" em número, para a planilha poder somar."""
    texto = str(valor).strip()
    if re.fullmatch(r"-?\d{1,3}(\.\d{3})*(,\d+)?|-?\d+(,\d+)?", texto):
        texto = texto.replace(".", "").replace(",", ".")
    try:
        numero = float(texto)
    except ValueError:
        return valor
    return int(numero) if numero.is_integer() and "." not in texto else numero


def _abrir_no_sistema(alvo: str) -> None:
    if sys.platform == "win32":
        os.startfile(alvo)  # noqa: S606 - só alvos conhecidos (programas da lista, arquivos, sites)
    elif sys.platform == "darwin":
        subprocess.Popen(["open", alvo])
    else:
        subprocess.Popen(["xdg-open", alvo])


def criar_ferramentas(
    confirmar: Callable[[str, str], bool],
    avisar: Callable[[str], None],
    pesquisar: Callable[[str], str] | None = None,
    ao_usar: Callable[[str], None] | None = None,
    analisar_imagem: Callable[[str, bytes], str] | None = None,
) -> list[Callable]:
    """Monta a lista de ferramentas.

    confirmar(titulo, detalhe) -> bool   pergunta ao usuário antes de ações sensíveis
    avisar(texto)                         mostra/fala um aviso (usado pelos lembretes)
    pesquisar(pergunta) -> str            pesquisa na internet (fornecido pelo cérebro)
    ao_usar(nome)                         avisa a tela qual ferramenta está em uso
    analisar_imagem(pergunta, png) -> str  descreve uma imagem (fornecido pelo cérebro)
    """
    PASTA_TRABALHO.mkdir(parents=True, exist_ok=True)
    arquivo_contatos = PASTA_TRABALHO / "contatos.json"

    def _contatos() -> dict:
        if arquivo_contatos.exists():
            return json.loads(arquivo_contatos.read_text(encoding="utf-8"))
        return {}

    # ---------- arquivos ----------

    def listar_arquivos(pasta: str = "") -> str:
        """Lista os arquivos de uma pasta. Sem pasta, lista a pasta de trabalho do Jarvis.

        Args:
          pasta: caminho da pasta (opcional).
        """
        alvo = _caminho(pasta) if pasta else PASTA_TRABALHO
        if not alvo.is_dir():
            return f"A pasta {alvo} não existe."
        itens = sorted(alvo.iterdir(), key=lambda p: (p.is_file(), p.name.lower()))[:200]
        linhas = [f"{'[pasta] ' if p.is_dir() else ''}{p.name}" for p in itens]
        return f"Conteúdo de {alvo}:\n" + ("\n".join(linhas) or "(vazia)")

    def ler_arquivo_texto(arquivo: str) -> str:
        """Lê um arquivo de texto (.txt, .csv, .py, .md, .json...).

        Args:
          arquivo: nome ou caminho do arquivo.
        """
        caminho = _caminho(arquivo)
        if not caminho.is_file():
            return f"Arquivo {caminho} não encontrado."
        return caminho.read_text(encoding="utf-8", errors="replace")[:20000]

    def salvar_arquivo_texto(arquivo: str, conteudo: str) -> str:
        """Salva um arquivo de texto: códigos (.py, .html, .js...), textos, relatórios, CSV.

        Args:
          arquivo: nome do arquivo com extensão, por exemplo "relatorio.txt" ou "robo.py".
          conteudo: texto completo do arquivo.
        """
        caminho = _caminho(arquivo)
        if caminho.exists() and not confirmar("Substituir arquivo?", f"O arquivo {caminho} já existe. Substituir?"):
            return "O usuário não autorizou substituir o arquivo."
        caminho.parent.mkdir(parents=True, exist_ok=True)
        caminho.write_text(conteudo, encoding="utf-8")
        return f"Arquivo salvo em {caminho}."

    def abrir(alvo: str) -> str:
        """Abre um arquivo, uma pasta, um site (http...) ou um programa do computador
        (excel, word, powerpoint, outlook, calculadora, bloco de notas, paint,
        explorador de arquivos, prompt de comando).

        Args:
          alvo: nome do programa, caminho do arquivo/pasta ou endereço do site.
        """
        nome = alvo.strip().lower()
        if nome in PROGRAMAS:
            if sys.platform != "win32":
                return "Abrir programas só funciona no Windows."
            _abrir_no_sistema(PROGRAMAS[nome])
            return f"Abrindo {nome}."
        if nome.startswith(("http://", "https://")):
            webbrowser.open(alvo.strip())
            return f"Abrindo {alvo}."
        caminho = _caminho(alvo) if alvo.strip() else PASTA_TRABALHO
        if not caminho.exists():
            return f"Não encontrei {caminho}. Programas disponíveis: {', '.join(PROGRAMAS)}."
        _abrir_no_sistema(str(caminho))
        return f"Abrindo {caminho}."

    # ---------- planilhas ----------

    def ler_planilha(arquivo: str, aba: str = "", max_linhas: int = 200) -> str:
        """Lê uma planilha Excel (.xlsx) e devolve o conteúdo, linha por linha.

        Args:
          arquivo: nome ou caminho da planilha.
          aba: nome da aba (opcional; sem ela, lê a aba ativa).
          max_linhas: quantas linhas ler no máximo.
        """
        from openpyxl import load_workbook

        caminho = _caminho(arquivo, ".xlsx")
        if not caminho.is_file():
            return f"Planilha {caminho} não encontrada."
        # Valores calculados pelo Excel; quando a fórmula ainda não foi calculada, mostra a fórmula.
        livro = load_workbook(caminho, read_only=True, data_only=True)
        formulas = load_workbook(caminho, read_only=True)
        planilha = livro[aba] if aba else livro.active
        planilha_formulas = formulas[planilha.title]
        linhas = []
        pares = zip(planilha.iter_rows(values_only=True), planilha_formulas.iter_rows(values_only=True))
        for numero, (valores, originais) in enumerate(pares, start=1):
            if numero > max_linhas:
                linhas.append(f"... (parei em {max_linhas} linhas)")
                break
            celulas = [o if v is None else v for v, o in zip(valores, originais)]
            linhas.append(f"{numero}: " + " | ".join("" if c is None else str(c) for c in celulas))
        livro.close()
        formulas.close()
        return (f"Planilha {caminho.name}, aba '{planilha.title}' "
                f"(abas: {', '.join(livro.sheetnames)}):\n" + "\n".join(linhas))

    def criar_planilha(arquivo: str, cabecalho: list[str], linhas: list[list[str]],
                       aba: str = "Planilha1") -> str:
        """Cria uma planilha Excel (.xlsx) nova com cabeçalho e dados.
        Números podem vir como texto ("1500", "1.234,56"): são convertidos para número.
        Fórmulas do Excel também funcionam (por exemplo "=SOMA(B2:B10)" deve ser escrito "=SUM(B2:B10)").

        Args:
          arquivo: nome da planilha, por exemplo "despesas.xlsx".
          cabecalho: títulos das colunas.
          linhas: lista de linhas, cada linha com os valores das colunas.
          aba: nome da aba.
        """
        from openpyxl import Workbook
        from openpyxl.styles import Font

        caminho = _caminho(arquivo, ".xlsx")
        if caminho.exists() and not confirmar("Substituir planilha?", f"A planilha {caminho} já existe. Substituir?"):
            return "O usuário não autorizou substituir a planilha."
        livro = Workbook()
        planilha = livro.active
        planilha.title = aba[:31]
        planilha.append(cabecalho)
        for celula in planilha[1]:
            celula.font = Font(bold=True)
        for linha in linhas:
            planilha.append([_numero_ou_texto(v) for v in linha])
        for coluna in planilha.columns:
            largura = max(len(str(c.value or "")) for c in coluna)
            planilha.column_dimensions[coluna[0].column_letter].width = min(max(10, largura + 2), 60)
        caminho.parent.mkdir(parents=True, exist_ok=True)
        livro.save(caminho)
        return f"Planilha criada em {caminho} com {len(linhas)} linhas."

    def adicionar_linhas_planilha(arquivo: str, linhas: list[list[str]], aba: str = "") -> str:
        """Acrescenta linhas no final de uma planilha Excel que já existe.

        Args:
          arquivo: nome ou caminho da planilha.
          linhas: linhas a acrescentar, cada uma com os valores das colunas.
          aba: nome da aba (opcional).
        """
        from openpyxl import load_workbook

        caminho = _caminho(arquivo, ".xlsx")
        if not caminho.is_file():
            return f"Planilha {caminho} não encontrada."
        if not confirmar("Alterar planilha?", f"Acrescentar {len(linhas)} linha(s) em {caminho.name}?"):
            return "O usuário não autorizou alterar a planilha."
        livro = load_workbook(caminho)
        planilha = livro[aba] if aba else livro.active
        for linha in linhas:
            planilha.append([_numero_ou_texto(v) for v in linha])
        try:
            livro.save(caminho)
        except PermissionError:
            return "Não consegui salvar: feche a planilha no Excel e peça de novo."
        return f"{len(linhas)} linha(s) acrescentadas em {caminho.name}."

    # ---------- código ----------

    def executar_codigo_python(codigo: str) -> str:
        """Executa um código Python no computador e devolve o que ele imprimir.
        Use para automações, cálculos, tratar dados ou gerar arquivos. O código roda na
        pasta de trabalho do Jarvis e pode usar a biblioteca openpyxl. O usuário
        precisa aprovar antes de rodar.

        Args:
          codigo: código Python completo.
        """
        if not confirmar("Executar código?", "O Jarvis quer executar este código Python:\n\n" + codigo[:1500]):
            return "O usuário não autorizou executar o código."
        interpretador = Path(sys.executable)
        if interpretador.name.lower() == "pythonw.exe":
            interpretador = interpretador.with_name("python.exe")
        try:
            resultado = subprocess.run(
                [str(interpretador), "-c", codigo], cwd=PASTA_TRABALHO, capture_output=True,
                text=True, encoding="utf-8", errors="replace", timeout=120,
                creationflags=getattr(subprocess, "CREATE_NO_WINDOW", 0),
            )
        except subprocess.TimeoutExpired:
            return "O código demorou mais de 2 minutos e foi interrompido."
        saida = (resultado.stdout + ("\nERROS:\n" + resultado.stderr if resultado.stderr else "")).strip()
        return f"Código terminou (código de saída {resultado.returncode}).\n{saida[-LIMITE_SAIDA:]}"

    # ---------- mensagens ----------

    def salvar_contato(nome: str, telefone: str) -> str:
        """Guarda um contato para mandar WhatsApp depois.

        Args:
          nome: nome do contato.
          telefone: telefone com DDD, por exemplo "83 99999-8888".
        """
        contatos = _contatos()
        contatos[nome.strip().lower()] = telefone
        arquivo_contatos.write_text(json.dumps(contatos, ensure_ascii=False, indent=2), encoding="utf-8")
        return f"Contato {nome} salvo."

    def preparar_whatsapp(contato: str, mensagem: str) -> str:
        """Abre o WhatsApp com a mensagem já escrita para o contato. O usuário só aperta enviar.

        Args:
          contato: nome de um contato salvo ou número com DDD.
          mensagem: texto da mensagem.
        """
        telefone = _contatos().get(contato.strip().lower(), contato)
        digitos = re.sub(r"\D", "", telefone)
        if len(digitos) < 10:
            return (f"Não tenho o número de {contato}. Peça o telefone ao usuário e use salvar_contato.")
        if len(digitos) <= 11:
            digitos = "55" + digitos  # Brasil
        webbrowser.open(f"https://wa.me/{digitos}?text={urllib.parse.quote(mensagem)}")
        return f"WhatsApp aberto com a mensagem para {contato}. Falta só o usuário apertar enviar."

    def preparar_email(para: str, assunto: str, corpo: str) -> str:
        """Abre um e-mail novo no Gmail já preenchido. O usuário revisa e aperta enviar.

        Args:
          para: e-mail do destinatário.
          assunto: assunto do e-mail.
          corpo: texto do e-mail.
        """
        parametros = urllib.parse.urlencode({"view": "cm", "to": para, "su": assunto, "body": corpo},
                                            quote_via=urllib.parse.quote)
        webbrowser.open(f"https://mail.google.com/mail/?{parametros}")
        return "E-mail aberto no Gmail, pronto para o usuário revisar e enviar."

    # ---------- lembretes ----------

    def criar_lembrete(minutos: int, texto: str) -> str:
        """Cria um lembrete: depois de alguns minutos o Jarvis avisa na tela e em voz alta.
        Funciona enquanto o Jarvis estiver aberto.

        Args:
          minutos: daqui a quantos minutos avisar.
          texto: o que lembrar.
        """
        temporizador = threading.Timer(max(0, minutos) * 60, avisar, args=(f"⏰ Lembrete: {texto}",))
        temporizador.daemon = True
        temporizador.start()
        return f"Lembrete criado para daqui a {minutos} minuto(s)."

    ferramentas = [listar_arquivos, ler_arquivo_texto, salvar_arquivo_texto, abrir,
                   ler_planilha, criar_planilha, adicionar_linhas_planilha,
                   executar_codigo_python, salvar_contato, preparar_whatsapp,
                   preparar_email, criar_lembrete]

    # ---------- memória ----------

    def lembrar_informacao(assunto: str, informacao: str) -> str:
        """Guarda uma informação para sempre (preferências, telefones, datas, nomes de clientes...).
        Use quando o usuário pedir para lembrar algo ou contar um fato pessoal importante.

        Args:
          assunto: título curto, por exemplo "aniversário da Yvnna".
          informacao: o que guardar.
        """
        return memoria.lembrar(assunto, informacao)

    def esquecer_informacao(assunto: str) -> str:
        """Apaga uma informação guardada na memória.

        Args:
          assunto: o título usado ao guardar.
        """
        return memoria.esquecer(assunto)

    ferramentas += [lembrar_informacao, esquecer_informacao]

    if analisar_imagem:
        def ver_tela(pergunta: str = "Descreva o que está na tela.") -> str:
            """Tira uma foto da tela do computador e responde sobre ela: ler um erro, resumir um
            documento aberto, explicar um gráfico, dizer o que está aberto.

            Args:
              pergunta: o que o usuário quer saber sobre a tela.
            """
            from PIL import ImageGrab

            imagem = ImageGrab.grab(all_screens=True)
            imagem.thumbnail((1920, 1920))
            buffer = io.BytesIO()
            imagem.save(buffer, format="PNG")
            return analisar_imagem(pergunta, buffer.getvalue())
        ferramentas.append(ver_tela)

    if pesquisar:
        def pesquisar_internet(pergunta: str) -> str:
            """Pesquisa no Google informações atuais: notícias, cotações, leis recentes,
            empresas, eventos. Devolve um resumo com as fontes.

            Args:
              pergunta: o que pesquisar.
            """
            return pesquisar(pergunta)
        ferramentas.append(pesquisar_internet)

    return [_protegida(f, ao_usar) for f in ferramentas]


def _protegida(funcao: Callable, ao_usar: Callable[[str], None] | None) -> Callable:
    """Avisa a tela qual ferramenta está rodando e transforma erros em texto para a IA."""
    @functools.wraps(funcao)
    def envolvida(*args, **kwargs):
        if ao_usar:
            ao_usar(funcao.__name__)
        try:
            return funcao(*args, **kwargs)
        except Exception as erro:
            return f"Erro em {funcao.__name__}: {type(erro).__name__}: {erro}"
    return envolvida
